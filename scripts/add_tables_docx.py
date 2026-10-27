import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc_path = r"C:\projects\news-aggregator\Mini_Project_Questionnaire(1).docx"
doc = docx.Document(doc_path)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), 'D3D3D3')
            borders.append(border)
        tblPr[0].append(borders)

# Find Q73 paragraph to insert Ablation Table after it
q73_idx = None
q74_idx = None

for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("73. If applicable"):
        q73_idx = i
    elif p.text.strip().startswith("74. What are the major limitations"):
        q74_idx = i

print(f"Q73 index: {q73_idx}, Q74 index: {q74_idx}")

# Check if tables already added
if len(doc.tables) <= 3:
    print("[*] Adding Table 3: Ablation Study Results...")
    # Add paragraph after Q73 answer
    ans_p = doc.paragraphs[q73_idx + 1]
    
    # Insert table
    table3 = doc.add_table(rows=5, cols=7)
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table3)
    
    t3_data = [
        ["Configuration", "Components Included", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"],
        ["Full Ensemble", "DistilBERT + TextCNN + BiLSTM (0.5/0.3/0.2)", "82.91%", "88.16%", "76.03%", "0.8164", "0.9361"],
        ["Ablation A: Drop DistilBERT", "TextCNN + BiLSTM (0.6/0.4)", "85.67%", "89.42%", "80.91%", "0.8495", "0.9292"],
        ["Ablation B: Drop TextCNN", "DistilBERT + BiLSTM (0.7/0.3)", "82.05%", "87.28%", "75.04%", "0.8070", "0.9024"],
        ["Ablation C: Drop BiLSTM", "DistilBERT + TextCNN (0.6/0.4)", "82.68%", "88.24%", "75.41%", "0.8132", "0.9403"]
    ]
    
    for r_idx, row in enumerate(table3.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t3_data[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 1 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.name = "Arial"
                if r_idx == 0:
                    run.font.bold = True
            if r_idx == 0:
                set_cell_background(cell, "F2F2F2")

    # Move table element right after the Q73 answer paragraph
    ans_p._p.addnext(table3._element)

    print("[*] Adding Table 4: Real News Headlines Domain-Shift Benchmark...")
    # Find Q74 answer paragraph
    ans_q74 = doc.paragraphs[q74_idx + 1]
    
    table4 = doc.add_table(rows=7, cols=4)
    table4.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table4)
    
    t4_data = [
        ["Model Architecture", "Headline Accuracy (60 Items)", "Correct / Total", "Domain Shift & Robustness Finding"],
        ["Logistic Regression (baseline)", "36.67%", "22 / 60", "Severe domain collapse; conflicting polarity on 'plot', 'role', 'support'"],
        ["Multinomial Naive Bayes (baseline)", "41.67%", "25 / 60", "Severe domain collapse; independence assumption fails on news syntax"],
        ["Bidirectional LSTM", "41.67%", "25 / 60", "Struggles with out-of-vocabulary news entities without pretrained embeddings"],
        ["TextCNN (1D-CNN)", "40.00%", "24 / 60", "Captures local n-grams but misclassifies neutral journalistic quotes"],
        ["DistilBERT (Transformer)", "70.00%", "42 / 60", "Strong semantic transfer due to deep bidirectional self-attention"],
        ["Deep Learning Ensemble (PulsePress)", "80.00%", "48 / 60", "Highest overall accuracy; combines contextual DL with neutrality zone"]
    ]
    
    for r_idx, row in enumerate(table4.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t4_data[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (1 <= c_idx <= 2) else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.name = "Arial"
                if r_idx == 0:
                    run.font.bold = True
            if r_idx == 0:
                set_cell_background(cell, "F2F2F2")
            elif r_idx == 6:
                set_cell_background(cell, "EBF5FB") # Soft blue highlight for proposed system

    # Move table element right after the Q74 answer paragraph
    ans_q74._p.addnext(table4._element)

doc.save(doc_path)
print(f"[+] Successfully inserted Ablation Study Table and Headline Benchmark Table into {doc_path}!")
